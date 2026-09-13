#!/usr/bin/env python3
"""Build the non-published governance database from the three source packages."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sqlite3
from collections.abc import Iterable, Iterator
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ASSET_VERSION = "1.0.0"
TIME_ZONE = "Asia/Shanghai"
BATCH_SIZE = 10_000
NULL_VALUES = {"", "NULL", "NONE", "N/A", "NA"}


SOURCE_SPECS = (
    ("metro_daily", "铁路和轨交", "铁路和轨交/REC_METRO_STN_FLOW_DAILY虹桥枢纽.csv", "PRODUCTION", True, None),
    ("train_production", "铁路和轨交", "铁路和轨交/TL_RT_TRAIN_LIST_VW20250920-.csv", "PRODUCTION", True, None),
    ("site_flow_may", "样例数据_五一前后", "样例数据_五一前后/14_V_ZHZX_SJSF_DYTRH_TRAFFIC_REC.xlsx", "SAMPLE", True, None),
    ("ridehail_match_may", "样例数据_五一前后", "样例数据_五一前后/19_V_ZHZX_YPSF_WYC_CM_INDICATOR_HQ_RT.xlsx", "SAMPLE", True, None),
    ("arrival_forecast_may", "样例数据_五一前后", "样例数据_五一前后/23_V_ZHZX_YPSF_DYTRH_ARRIVAL_PF_LT_RT.xlsx", "SAMPLE", True, None),
    ("evac_ratio_hour_may", "样例数据_五一前后", "样例数据_五一前后/24_V_ZHZX_YPSF_DYTRH_EVAC_MODE_RATIO_LT_RT.xlsx", "SAMPLE", True, None),
    ("evac_ratio_day_may", "样例数据_五一前后", "样例数据_五一前后/25_V_ZHZX_YPSF_DYTRH_EVAC_MODE_RATIO_DAY_RT.xlsx", "SAMPLE", True, None),
    ("train_may", "样例数据_五一前后", "样例数据_五一前后/4_TL_RT_TRAIN_LIST_VW.xlsx", "SAMPLE", True, None),
    ("ridehail_demand", "20250920-1021", "20250920-1021/SHDATA.T_DSJZX_WYCXT_WYC_DQYCXQ_RT.xlsx", "PRODUCTION", True, None),
    ("ridehail_orders", "20250920-1021", "20250920-1021/V_DSJZX_WYCXT_WYC_ZBDL_RT.xlsx", "PRODUCTION", True, None),
    ("taxi_yard", "20250920-1021", "20250920-1021/V_ZHZX_RHPT_XYC_CURRENT_NUMBER_RT.xlsx", "PRODUCTION", True, None),
    ("taxi_yard_duplicate", "20250920-1021", "20250920-1021/V_ZHZX_RHPT_XYC_CURRENT_NUMBER_RT_VW.xlsx", "DUPLICATE_COPY", False, "taxi_yard"),
    ("site_flow_national_day", "20250920-1021", "20250920-1021/V_ZHZX_SJSF_DYTRH_FIFTEENNUM_RT.xlsx", "PRODUCTION", True, None),
    ("arrival_forecast_national_day", "20250920-1021", "20250920-1021/V_ZHZX_YPSF_DYTRH_ARRIVAL_PF_LT_RT.xlsx", "PRODUCTION", True, None),
    ("evac_ratio_hour_national_day", "20250920-1021", "20250920-1021/V_ZHZX_YPSF_DYTRH_EVAC_MODE_RATIO_LT_RT.xlsx", "PRODUCTION", True, None),
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_header(value: object) -> str:
    return str(value).strip().lstrip("?").strip('"').strip()


def nullish(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip().upper() in NULL_VALUES)


def text(value: object, *, trim: bool = True) -> str | None:
    if nullish(value):
        return None
    result = str(value)
    return result.strip() if trim else result


def parse_datetime(value: object) -> datetime | None:
    if nullish(value):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    raw = str(value).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y/%m/%d",
    ):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    raise ValueError(f"Unsupported datetime: {raw}")


def iso_datetime(value: object) -> str | None:
    parsed = parse_datetime(value)
    if parsed is None:
        return None
    return parsed.isoformat(sep=" ", timespec="microseconds" if parsed.microsecond else "seconds")


def iso_date(value: object) -> str | None:
    parsed = parse_datetime(value)
    return parsed.date().isoformat() if parsed else None


def integer(value: object) -> int | None:
    if nullish(value):
        return None
    return int(float(str(value).strip()))


def real(value: object) -> float | None:
    if nullish(value):
        return None
    return float(str(value).strip())


def record_hash(values: Iterable[object]) -> str:
    payload = json.dumps(list(values), ensure_ascii=False, separators=(",", ":"), default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def iter_csv(path: Path) -> Iterator[tuple[list[str], int, tuple[object, ...]]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle)
        headers = [clean_header(value) for value in next(reader)]
        for row_number, row in enumerate(reader, start=2):
            yield headers, row_number, tuple(row)


def iter_xlsx(path: Path) -> Iterator[tuple[list[str], int, tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    if worksheet.max_row == 1 and worksheet.max_column == 1:
        worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    headers = [clean_header(value) for value in next(rows)]
    try:
        for row_number, row in enumerate(rows, start=2):
            if any(not nullish(value) for value in row):
                yield headers, row_number, tuple(row)
    finally:
        workbook.close()


def iter_rows(path: Path) -> Iterator[tuple[list[str], int, tuple[object, ...]]]:
    return iter_csv(path) if path.suffix.lower() == ".csv" else iter_xlsx(path)


def extract_condition(path: Path) -> str | None:
    if path.suffix.lower() != ".xlsx":
        return None
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) < 2:
            return None
        value = workbook.worksheets[1].cell(1, 1).value
        return text(value)
    finally:
        workbook.close()


def row_dict(headers: list[str], values: tuple[object, ...]) -> dict[str, object]:
    return {name: values[index] if index < len(values) else None for index, name in enumerate(headers)}


def batches(rows: Iterable[tuple[Any, ...]], size: int = BATCH_SIZE) -> Iterator[list[tuple[Any, ...]]]:
    batch: list[tuple[Any, ...]] = []
    for row in rows:
        batch.append(row)
        if len(batch) >= size:
            yield batch
            batch = []
    if batch:
        yield batch


def create_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA journal_mode=DELETE;
        PRAGMA synchronous=NORMAL;
        PRAGMA foreign_keys=ON;
        PRAGMA temp_store=MEMORY;
        PRAGMA cache_size=-200000;

        CREATE TABLE meta_build (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        ) STRICT;

        CREATE TABLE meta_source_file (
            source_file_id TEXT PRIMARY KEY,
            package_name TEXT NOT NULL,
            relative_path TEXT NOT NULL UNIQUE,
            file_sha256 TEXT NOT NULL,
            size_bytes INTEGER NOT NULL,
            dataset_role TEXT NOT NULL,
            is_canonical INTEGER NOT NULL CHECK (is_canonical IN (0, 1)),
            duplicate_of TEXT REFERENCES meta_source_file(source_file_id),
            source_row_count INTEGER NOT NULL DEFAULT 0,
            accepted_row_count INTEGER NOT NULL DEFAULT 0,
            source_column_count INTEGER,
            source_sheet TEXT,
            extract_condition TEXT
        ) STRICT;

        CREATE TABLE meta_quality_issue (
            issue_code TEXT PRIMARY KEY,
            severity TEXT NOT NULL CHECK (severity IN ('ERROR', 'WARN')),
            domain_code TEXT NOT NULL,
            observed_value TEXT NOT NULL,
            description_cn TEXT NOT NULL
        ) STRICT;

        CREATE TABLE meta_conflict_summary (
            conflict_code TEXT PRIMARY KEY,
            record_count INTEGER NOT NULL,
            resolution TEXT NOT NULL,
            description_cn TEXT NOT NULL
        ) STRICT;

        CREATE TABLE raw_train_stop (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL REFERENCES meta_source_file(source_file_id),
            source_row_number INTEGER NOT NULL,
            dataset_role TEXT NOT NULL,
            departure_date TEXT NOT NULL,
            train_code TEXT NOT NULL,
            start_station_name TEXT,
            end_station_name TEXT,
            start_time TEXT,
            end_time TEXT,
            train_type TEXT,
            station_code TEXT NOT NULL,
            station_name TEXT,
            station_train_code TEXT,
            fact_reach_time TEXT,
            def_reach_time TEXT,
            fact_leave_time TEXT,
            def_leave_time TEXT,
            reach_time TEXT,
            leave_time TEXT,
            late_time TEXT,
            diff_days INTEGER,
            pre_get_off INTEGER,
            pre_get_on INTEGER,
            actual_get_off INTEGER,
            actual_get_on INTEGER,
            get_on INTEGER,
            get_off INTEGER,
            late_flag INTEGER,
            stop_flag INTEGER,
            del_flag INTEGER,
            rg_del_flag INTEGER,
            pre_update_time TEXT,
            fact_update_time TEXT,
            update_time TEXT,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;

        CREATE TABLE raw_metro_flow_day (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL,
            source_row_number INTEGER NOT NULL,
            line_id TEXT NOT NULL,
            station_id TEXT NOT NULL,
            station_name TEXT NOT NULL,
            service_date TEXT NOT NULL,
            inbound_flow INTEGER NOT NULL,
            outbound_flow INTEGER NOT NULL,
            total_flow INTEGER NOT NULL,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;

        CREATE TABLE raw_site_flow_observation (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL,
            source_row_number INTEGER NOT NULL,
            dataset_role TEXT NOT NULL,
            observed_at TEXT NOT NULL,
            funsite_code TEXT NOT NULL,
            funsite_name TEXT,
            function_code TEXT,
            hub_code_raw TEXT,
            hub_code TEXT,
            measurement_point_code TEXT,
            metric_code TEXT NOT NULL,
            observed_value REAL,
            updated_at TEXT,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;

        CREATE TABLE raw_arrival_forecast (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL,
            source_row_number INTEGER NOT NULL,
            dataset_role TEXT NOT NULL,
            hub_code TEXT NOT NULL,
            hub_name TEXT,
            predicted_date TEXT NOT NULL,
            predicted_at TEXT NOT NULL,
            predicted_passengers INTEGER,
            issued_at TEXT NOT NULL,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;

        CREATE TABLE raw_evac_ratio_hour (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL,
            source_row_number INTEGER NOT NULL,
            dataset_role TEXT NOT NULL,
            hub_code TEXT NOT NULL,
            hub_name TEXT,
            predicted_date TEXT NOT NULL,
            predicted_at TEXT NOT NULL,
            metro_ratio REAL,
            bus_ratio REAL,
            taxi_ratio REAL,
            ridehail_ratio REAL,
            car_ratio REAL,
            issued_at TEXT NOT NULL,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;

        CREATE TABLE raw_evac_ratio_day (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL,
            source_row_number INTEGER NOT NULL,
            dataset_role TEXT NOT NULL,
            issued_at TEXT NOT NULL,
            hub_code TEXT NOT NULL,
            hub_name TEXT,
            predicted_date TEXT NOT NULL,
            metro_ratio REAL,
            bus_ratio REAL,
            taxi_ratio REAL,
            ridehail_ratio REAL,
            car_ratio REAL,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;

        CREATE TABLE raw_taxi_yard_state (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL,
            source_row_number INTEGER NOT NULL,
            observed_at TEXT NOT NULL,
            yard_name TEXT NOT NULL,
            max_vehicle_count INTEGER,
            inbound_vehicle_count INTEGER,
            outbound_vehicle_count INTEGER,
            current_vehicle_count INTEGER,
            avg_waiting_minutes REAL,
            pjyl REAL,
            traffic_15sum REAL,
            ratio REAL,
            car_15sum REAL,
            status_code TEXT,
            qbb_status_code TEXT,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;

        CREATE TABLE raw_ridehail_demand (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL,
            source_row_number INTEGER NOT NULL,
            observed_at TEXT NOT NULL,
            hub_code TEXT NOT NULL,
            hub_name TEXT,
            demand_count INTEGER,
            status_code TEXT,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;

        CREATE TABLE raw_ridehail_order_window (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL,
            source_row_number INTEGER NOT NULL,
            observed_at TEXT NOT NULL,
            hub_code TEXT NOT NULL,
            hub_name TEXT,
            orders_3min INTEGER,
            orders_5min INTEGER,
            orders_10min INTEGER,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;

        CREATE TABLE raw_ridehail_match (
            record_key INTEGER PRIMARY KEY,
            source_file_id TEXT NOT NULL,
            source_row_number INTEGER NOT NULL,
            dataset_role TEXT NOT NULL,
            observed_at TEXT NOT NULL,
            sample_start_at TEXT,
            sample_end_at TEXT,
            created_orders INTEGER,
            identified_orders INTEGER,
            matched_orders INTEGER,
            match_ratio REAL,
            create_to_match_avg_minutes REAL,
            create_to_match_median_minutes REAL,
            status_code TEXT,
            source_record_hash TEXT NOT NULL,
            UNIQUE (source_file_id, source_row_number)
        ) STRICT;
        """
    )


def transform_train(file_id: str, role: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        normalized = (
            file_id, row_number, role, iso_date(r.get("DEPARTURE_DATE")), text(r.get("TRAIN_CODE")),
            text(r.get("START_STATION_NAME")), text(r.get("END_STATION_NAME")), iso_datetime(r.get("START_TIME")),
            iso_datetime(r.get("END_TIME")), text(r.get("TYPE")), text(r.get("STATION_CODE")),
            text(r.get("STATION_NAME")), text(r.get("STATION_TRAIN_CODE")), iso_datetime(r.get("FACT_REACH_TIME")),
            iso_datetime(r.get("DEF_REACH_TIME")), iso_datetime(r.get("FACT_LEAVE_TIME")), iso_datetime(r.get("DEF_LEAVE_TIME")),
            iso_datetime(r.get("REACH_TIME")), iso_datetime(r.get("LEAVE_TIME")), iso_datetime(r.get("LATE_TIME")),
            integer(r.get("DIFF_DAYS")), integer(r.get("PRE_GET_OFF")), integer(r.get("PRE_GET_ON")),
            integer(r.get("ACTUAL_GET_OFF")), integer(r.get("ACTUAL_GET_ON")), integer(r.get("GET_ON")),
            integer(r.get("GET_OFF")), integer(r.get("LATE_FLAG")), integer(r.get("STOP_FLAG")), integer(r.get("DEL_FLAG")),
            integer(r.get("RG_DEL_FLAG")), iso_datetime(r.get("PRE_UPDATE_TIME")), iso_datetime(r.get("FACT_UPDATE_TIME")),
            iso_datetime(r.get("UPDATE_TIME")),
        )
        yield (*normalized, record_hash(normalized[3:]))


def transform_metro(file_id: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        normalized = (
            file_id, row_number, text(r["LINE_ID"]), text(r["STATN_ID"]), text(r["STATN_NM_CN"]),
            iso_date(r["FDT_DATE"]), integer(r["IN_FLOW"]), integer(r["OUT_FLOW"]), integer(r["DAILY_FLOW"]),
        )
        yield (*normalized, record_hash(normalized[2:]))


def transform_site_flow(file_id: str, role: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        raw_hub = text(r.get("FSTR_CZ_CODE"), trim=False)
        metric = "TRAFFICNUM" if "TRAFFICNUM" in r else "INNUM"
        normalized = (
            file_id, row_number, role, iso_datetime(r["DATA_TIME"]), text(r["FSTR_FUNSITE_CODE"]),
            text(r.get("FSTR_FUNSITE_NAME")), text(r.get("FSTR_FUN_CODE")), raw_hub, text(raw_hub),
            text(r.get("FSTR_MP_CODE")), metric, real(r.get(metric)), iso_datetime(r.get("UPDATE_TIME")),
        )
        yield (*normalized, record_hash(normalized[3:]))


def transform_arrival(file_id: str, role: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        normalized = (
            file_id, row_number, role, text(r["HUB_CODE"]), text(r.get("HUB_NAME")), iso_date(r["PRED_DATE"]),
            iso_datetime(r["PRED_TIME"]), integer(r.get("PRED_PF")), iso_datetime(r["UPDATE_TIME"]),
        )
        yield (*normalized, record_hash(normalized[3:]))


def transform_evac_hour(file_id: str, role: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        normalized = (
            file_id, row_number, role, text(r["HUB_CODE"]), text(r.get("HUB_NAME")), iso_date(r["PRED_DATE"]),
            iso_datetime(r["PRED_TIME"]), real(r.get("METRO_RATIO")), real(r.get("BUS_RATIO")), real(r.get("TAXI_RATIO")),
            real(r.get("NETCAR_RATIO")), real(r.get("CAR_RATIO")), iso_datetime(r["UPDATE_TIME"]),
        )
        yield (*normalized, record_hash(normalized[3:]))


def transform_evac_day(file_id: str, role: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        normalized = (
            file_id, row_number, role, iso_datetime(r["UPDATE_TIME"]), text(r["HUB_CODE"]), text(r.get("HUB_NAME")),
            iso_date(r["PRED_DATE"]), real(r.get("METRO_RATIO")), real(r.get("BUS_RATIO")), real(r.get("TAXI_RATIO")),
            real(r.get("NETCAR_RATIO")), real(r.get("CAR_RATIO")),
        )
        yield (*normalized, record_hash(normalized[3:]))


def transform_taxi_yard(file_id: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        normalized = (
            file_id, row_number, iso_datetime(r["DATATIME"]), text(r["FSTR_FUNSITE_NAME"]), integer(r.get("MAX_CAR_NUM")),
            integer(r.get("INCOUNT_NUM")), integer(r.get("OUTCOUNT_NUM")), integer(r.get("CURRENT_VEHICLE_NUM")),
            real(r.get("AVG_WAITING_TIME")), real(r.get("PJYL")), real(r.get("TRAFFICNUM_15SUM")), real(r.get("RATIO")),
            real(r.get("CAR_15SUM")), text(r.get("FSTR_ZT")), text(r.get("FSTR_QBB_ZT")),
        )
        yield (*normalized, record_hash(normalized[2:]))


def transform_demand(file_id: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        normalized = (
            file_id, row_number, iso_datetime(r["UPDATE_TIME"]), text(r["FSTR_CZ_CODE"]), text(r.get("FSTR_CZ_NAME")),
            integer(r.get("CNT")), text(r.get("ZT")),
        )
        yield (*normalized, record_hash(normalized[2:]))


def transform_orders(file_id: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        normalized = (
            file_id, row_number, iso_datetime(r["UPDATE_TIME"]), text(r["FSTR_CZ_CODE"]), text(r.get("FSTR_CZ_NAME")),
            integer(r.get("THREE")), integer(r.get("FIVE")), integer(r.get("TEN")),
        )
        yield (*normalized, record_hash(normalized[2:]))


def transform_match(file_id: str, role: str, path: Path) -> Iterator[tuple[Any, ...]]:
    for headers, row_number, values in iter_rows(path):
        r = row_dict(headers, values)
        normalized = (
            file_id, row_number, role, iso_datetime(r["UPDATE_TIME"]), iso_datetime(r.get("SAMPLE_START_TIME")),
            iso_datetime(r.get("SAMPLE_END_TIME")), integer(r.get("CREATED_ORDERS")), integer(r.get("IDENTIFIED_ORDERS")),
            integer(r.get("MATCH_ORDERS")), real(r.get("MATCH_RATIO")), real(r.get("CREATE_TO_MATCH_MIN")),
            real(r.get("CREATE_TO_MATCH_MEDIAN_MIN")), text(r.get("STATUS")),
        )
        yield (*normalized, record_hash(normalized[3:]))


IMPORTERS = {
    "metro_daily": ("raw_metro_flow_day", transform_metro),
    "train_production": ("raw_train_stop", transform_train),
    "train_may": ("raw_train_stop", transform_train),
    "site_flow_may": ("raw_site_flow_observation", transform_site_flow),
    "site_flow_national_day": ("raw_site_flow_observation", transform_site_flow),
    "ridehail_match_may": ("raw_ridehail_match", transform_match),
    "arrival_forecast_may": ("raw_arrival_forecast", transform_arrival),
    "arrival_forecast_national_day": ("raw_arrival_forecast", transform_arrival),
    "evac_ratio_hour_may": ("raw_evac_ratio_hour", transform_evac_hour),
    "evac_ratio_hour_national_day": ("raw_evac_ratio_hour", transform_evac_hour),
    "evac_ratio_day_may": ("raw_evac_ratio_day", transform_evac_day),
    "taxi_yard": ("raw_taxi_yard_state", transform_taxi_yard),
    "ridehail_demand": ("raw_ridehail_demand", transform_demand),
    "ridehail_orders": ("raw_ridehail_order_window", transform_orders),
}


def insert_rows(conn: sqlite3.Connection, table: str, rows: Iterable[tuple[Any, ...]]) -> int:
    column_count = conn.execute(f"SELECT COUNT(*) FROM pragma_table_info('{table}') WHERE name <> 'record_key'").fetchone()[0]
    sql = f"INSERT INTO {table} VALUES (NULL, {','.join('?' for _ in range(column_count))})"
    count = 0
    for batch in batches(rows):
        conn.executemany(sql, batch)
        count += len(batch)
    return count


def source_shape(path: Path) -> tuple[int, int]:
    rows = 0
    columns = 0
    for headers, _, _ in iter_rows(path):
        columns = len(headers)
        rows += 1
    return rows, columns


def normalized_xlsx_digest(path: Path) -> tuple[int, str]:
    digest = hashlib.sha256()
    count = 0
    for headers, _, values in iter_xlsx(path):
        normalized = [None if nullish(value) else str(value).strip() for value in values[: len(headers)]]
        digest.update(json.dumps(normalized, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
        count += 1
    return count, digest.hexdigest()


def register_sources(conn: sqlite3.Connection, source_root: Path) -> dict[str, Path]:
    paths: dict[str, Path] = {}
    for file_id, package, relative, role, canonical, duplicate_of in SOURCE_SPECS:
        path = source_root / relative
        if not path.is_file():
            raise FileNotFoundError(path)
        paths[file_id] = path
        conn.execute(
            "INSERT INTO meta_source_file "
            "(source_file_id, package_name, relative_path, file_sha256, size_bytes, dataset_role, is_canonical, duplicate_of, source_sheet, extract_condition) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (file_id, package, relative, sha256_file(path), path.stat().st_size, role, int(canonical), duplicate_of, "Sheet1/结果集" if path.suffix.lower() == ".xlsx" else None, extract_condition(path)),
        )
    conn.commit()
    return paths


def import_sources(conn: sqlite3.Connection, paths: dict[str, Path]) -> None:
    roles = {spec[0]: spec[3] for spec in SOURCE_SPECS}
    for file_id, _, _, _, canonical, _ in SOURCE_SPECS:
        path = paths[file_id]
        if not canonical:
            row_count, column_count = source_shape(path)
            conn.execute(
                "UPDATE meta_source_file SET source_row_count=?, source_column_count=? WHERE source_file_id=?",
                (row_count, column_count, file_id),
            )
            conn.commit()
            print(f"Registered duplicate {file_id}: {row_count:,} rows", flush=True)
            continue
        table, importer = IMPORTERS[file_id]
        role = roles[file_id]
        transformed = importer(file_id, role, path) if importer in (transform_train, transform_site_flow, transform_match, transform_arrival, transform_evac_hour, transform_evac_day) else importer(file_id, path)
        count = insert_rows(conn, table, transformed)
        headers = next(iter_rows(path))[0]
        conn.execute(
            "UPDATE meta_source_file SET source_row_count=?, accepted_row_count=?, source_column_count=? WHERE source_file_id=?",
            (count, count, len(headers), file_id),
        )
        conn.commit()
        print(f"Imported {file_id}: {count:,} rows into {table}", flush=True)


def build_governance_results(conn: sqlite3.Connection, paths: dict[str, Path]) -> None:
    conn.execute("DELETE FROM meta_conflict_summary")
    conn.execute("DELETE FROM meta_quality_issue")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_train_business_key ON raw_train_stop(departure_date, train_code, station_code)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_train_station_time ON raw_train_stop(station_code, reach_time, leave_time)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_metro_date_station ON raw_metro_flow_day(service_date, station_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_site_time ON raw_site_flow_observation(funsite_code, observed_at)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_arrival_key ON raw_arrival_forecast(hub_code, predicted_at, issued_at)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_evac_key ON raw_evac_ratio_hour(hub_code, predicted_at, issued_at)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_taxi_time ON raw_taxi_yard_state(yard_name, observed_at)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_demand_time ON raw_ridehail_demand(hub_code, observed_at)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_orders_time ON raw_ridehail_order_window(hub_code, observed_at)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_raw_match_time ON raw_ridehail_match(observed_at)")
    conn.commit()

    canonical_count, canonical_digest = normalized_xlsx_digest(paths["taxi_yard"])
    duplicate_count, duplicate_digest = normalized_xlsx_digest(paths["taxi_yard_duplicate"])
    if (canonical_count, canonical_digest) != (duplicate_count, duplicate_digest):
        raise RuntimeError("Taxi yard RT and RT_VW are not exact normalized duplicates")

    overlap = conn.execute(
        "SELECT COUNT(*) FROM raw_train_stop s JOIN raw_train_stop p "
        "ON p.dataset_role='PRODUCTION' AND s.dataset_role='SAMPLE' "
        "AND p.departure_date=s.departure_date AND p.train_code=s.train_code AND p.station_code=s.station_code"
    ).fetchone()[0]
    sample_only = conn.execute(
        "SELECT COUNT(*) FROM raw_train_stop s WHERE s.dataset_role='SAMPLE' AND NOT EXISTS ("
        "SELECT 1 FROM raw_train_stop p WHERE p.dataset_role='PRODUCTION' AND p.departure_date=s.departure_date "
        "AND p.train_code=s.train_code AND p.station_code=s.station_code)"
    ).fetchone()[0]
    late_conflict = conn.execute(
        "SELECT COUNT(*) FROM raw_train_stop s JOIN raw_train_stop p "
        "ON p.dataset_role='PRODUCTION' AND s.dataset_role='SAMPLE' "
        "AND p.departure_date=s.departure_date AND p.train_code=s.train_code AND p.station_code=s.station_code "
        "WHERE COALESCE(p.late_time,'')<>COALESCE(s.late_time,'') OR COALESCE(p.late_flag,-1)<>COALESCE(s.late_flag,-1)"
    ).fetchone()[0]
    conn.executemany(
        "INSERT INTO meta_conflict_summary VALUES (?, ?, ?, ?)",
        (
            ("TAXI_YARD_EXACT_DUPLICATE", duplicate_count, "发布 taxi_yard，副本只保留文件血缘", "RT 与 RT_VW 规范化内容完全一致"),
            ("TRAIN_SAMPLE_OVERLAP", overlap, "默认事实选择 PRODUCTION", "五一铁路样例与全年铁路表业务键重合"),
            ("TRAIN_SAMPLE_ONLY", sample_only, "保留在 std，默认事实排除", "只存在于样例的铁路业务键"),
            ("TRAIN_LATE_FIELD_CONFLICT", late_conflict, "保留两版，默认事实选择 PRODUCTION", "重合铁路记录的晚点时间或标志不同"),
        ),
    )

    issues = (
        ("METRO_MISSING_GRID", "WARN", "RAIL", "6", "2026-04-08 和 04-09 缺少 6 个日期—线路—站点组合"),
        ("TAXI_YARD_NULL_FIELDS", "WARN", "HUBOPS", str(conn.execute("SELECT COUNT(*) FROM raw_taxi_yard_state WHERE max_vehicle_count IS NULL OR inbound_vehicle_count IS NULL OR outbound_vehicle_count IS NULL OR avg_waiting_minutes IS NULL").fetchone()[0]), "出租车场容量、进出或等客字段缺失记录"),
        ("MATCH_NEGATIVE_DURATION", "WARN", "RIDEHAIL", str(conn.execute("SELECT COUNT(*) FROM raw_ridehail_match WHERE create_to_match_avg_minutes<0 OR create_to_match_median_minutes<0").fetchone()[0]), "撮合耗时为负的记录"),
        ("EVAC_INCOMPLETE", "WARN", "FORECAST", str(conn.execute("SELECT COUNT(*) FROM raw_evac_ratio_hour WHERE metro_ratio IS NULL OR bus_ratio IS NULL OR taxi_ratio IS NULL OR ridehail_ratio IS NULL OR car_ratio IS NULL").fetchone()[0]), "小时疏散比例五种方式不完整的记录"),
        ("EVAC_RATIO_SUM_GT_ONE", "WARN", "FORECAST", str(conn.execute("SELECT COUNT(*) FROM raw_evac_ratio_hour WHERE COALESCE(metro_ratio,0)+COALESCE(bus_ratio,0)+COALESCE(taxi_ratio,0)+COALESCE(ridehail_ratio,0)+COALESCE(car_ratio,0)>1.000001").fetchone()[0]), "小时疏散比例合计大于 1 的记录"),
        ("SITE_FLOW_MIXED_INTERVAL", "WARN", "HUBOPS", "15/1/3 minutes", "功能点流量存在多种主要落库间隔，指标窗口未确认"),
        ("STATUS_DICTIONARY_MISSING", "WARN", "RIDEHAIL", "ZT/FSTR_ZT/FSTR_QBB_ZT/STATUS", "状态码业务含义未确认"),
        ("PASSENGER_TAXI_WAIT_UNAVAILABLE", "WARN", "HUBOPS", "unavailable", "现有数据没有旅客出租车等待分钟数"),
    )
    conn.executemany("INSERT INTO meta_quality_issue VALUES (?, ?, ?, ?, ?)", issues)
    conn.commit()


def validate_governance(conn: sqlite3.Connection) -> None:
    physical = conn.execute("SELECT SUM(source_row_count) FROM meta_source_file").fetchone()[0]
    accepted = conn.execute("SELECT SUM(accepted_row_count) FROM meta_source_file").fetchone()[0]
    if physical != 3_602_908:
        raise RuntimeError(f"Unexpected physical row count: {physical}")
    if accepted != 3_594_268:
        raise RuntimeError(f"Unexpected accepted row count: {accepted}")
    expected = {
        "raw_train_stop": 801_433,
        "raw_metro_flow_day": 2_388,
        "raw_site_flow_observation": 1_236_100,
        "raw_ridehail_match": 11_148,
        "raw_arrival_forecast": 750_912,
        "raw_evac_ratio_hour": 750_912,
        "raw_evac_ratio_day": 40,
        "raw_taxi_yard_state": 8_640,
        "raw_ridehail_demand": 17_833,
        "raw_ridehail_order_window": 14_862,
    }
    for table, expected_count in expected.items():
        actual = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        if actual != expected_count:
            raise RuntimeError(f"{table}: expected {expected_count}, got {actual}")
    duplicate_keys = conn.execute(
        "SELECT COUNT(*) FROM (SELECT departure_date,train_code,station_code,COUNT(*) n FROM raw_train_stop "
        "WHERE dataset_role='PRODUCTION' GROUP BY 1,2,3 HAVING n>1)"
    ).fetchone()[0]
    if duplicate_keys:
        raise RuntimeError(f"Production train business-key duplicates: {duplicate_keys}")
    if conn.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
        raise RuntimeError("Governance database integrity check failed")


def build(source_root: Path, output: Path) -> None:
    source_root = source_root.resolve()
    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    for suffix in ("", "-wal", "-shm"):
        candidate = Path(f"{output}{suffix}")
        if candidate.exists():
            candidate.unlink()
    conn = sqlite3.connect(output)
    create_schema(conn)
    conn.executemany(
        "INSERT INTO meta_build VALUES (?, ?)",
        (
            ("asset_version", ASSET_VERSION),
            ("built_at", datetime.now().astimezone().isoformat(timespec="seconds")),
            ("canonical_time_zone", TIME_ZONE),
            ("source_root_sha_policy", "SHA-256 per source file"),
            ("physical_source_row_count", "3602908"),
            ("canonical_accepted_row_count", "3594268"),
        ),
    )
    paths = register_sources(conn, source_root)
    import_sources(conn, paths)
    build_governance_results(conn, paths)
    validate_governance(conn)
    conn.execute("PRAGMA optimize")
    conn.close()
    print(f"Built governance database: {output}", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the Shanghai hub traffic governance SQLite database.")
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source_root, args.output)


if __name__ == "__main__":
    main()
